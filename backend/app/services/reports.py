import io
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas
from typing import List, Any

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.HexColor("#64748b"))
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(750, 25, page_text)
        self.drawString(40, 25, "Rajiv Gandhi University • Faculty Activity Management Portal • Confidential")
        self.restoreState()

def generate_pdf_report(
    title: str,
    subtitle: str,
    headers: List[str],
    data: List[List[Any]],
    col_widths: List[float] = None
) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=36,
        rightMargin=36,
        topMargin=40,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1e293b"),
        alignment=0,
        spaceAfter=4
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#64748b"),
        spaceAfter=12
    )

    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#334155")
    )

    header_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=11,
        textColor=colors.white
    )

    elements = []
    
    elements.append(Paragraph("Rajiv Gandhi University", ParagraphStyle(
        'InstName', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=12, textColor=colors.HexColor("#4338ca")
    )))
    elements.append(Paragraph(title, title_style))
    gen_time = datetime.datetime.now().strftime("%B %d, %Y - %I:%M %p")
    elements.append(Paragraph(f"{subtitle} | Generated on: {gen_time}", subtitle_style))
    elements.append(Spacer(1, 8))

    formatted_headers = [Paragraph(f"<b>{h}</b>", header_style) for h in headers]
    formatted_rows = [formatted_headers]

    for row in data:
        formatted_row = []
        for item in row:
            text = str(item) if item is not None else "-"
            if len(text) > 200:
                text = text[:197] + "..."
            formatted_row.append(Paragraph(text, cell_style))
        formatted_rows.append(formatted_row)

    if not col_widths:
        avail_width = 720
        col_widths = [avail_width / len(headers)] * len(headers)

    table = Table(formatted_rows, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#312e81")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ('TOPPADDING', (0, 1), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
    ]))

    elements.append(table)
    doc.build(elements, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer

def generate_excel_report(
    title: str,
    headers: List[str],
    data: List[List[Any]],
    sheet_name: str = "Report"
) -> io.BytesIO:
    buffer = io.BytesIO()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:30]

    # Title row
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 4))
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = f"Rajiv Gandhi University - {title}"
    title_cell.font = Font(name='Calibri', size=16, bold=True, color='1E293B')
    
    gen_cell = ws.cell(row=2, column=1)
    gen_cell.value = f"Generated on {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    gen_cell.font = Font(name='Calibri', size=10, italic=True, color='64748B')

    # Header styling
    header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    header_font = Font(name='Calibri', size=11, bold=True, color="FFFFFF")
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data styling
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for r_idx, row in enumerate(data, start=5):
        is_alt = (r_idx % 2 == 0)
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.value = val
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
            if is_alt:
                cell.fill = alt_fill

    # Auto-fit columns
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val = str(cell.value or '')
            if cell.row > 2:
                max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

    wb.save(buffer)
    buffer.seek(0)
    return buffer
