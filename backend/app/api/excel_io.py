from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import Optional, List, Dict, Any
import openpyxl
import io
import datetime
from app.database import get_db
from app.models import models
from app.api.auth import get_current_user
from app.services.audit import log_action

router = APIRouter()

def get_user_faculty(user: models.User, db: Session):
    if user.faculty_profile:
        return user.faculty_profile
    return db.query(models.Faculty).filter(models.Faculty.email == user.email).first()

def parse_date_safely(val):
    if val is None:
        return None
    if isinstance(val, (datetime.date, datetime.datetime)):
        return val.date() if isinstance(val, datetime.datetime) else val
        
    val_str = str(val).strip()
    if val_str.lower() in ["dd/mm/yyyy", "dd-mm-yyyy", "yyyy-mm-dd", "-", "", "none", "null"]:
        return None
        
    for fmt in ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%Y/%m/%d"]:
        try:
            return datetime.datetime.strptime(val_str, fmt).date()
        except ValueError:
            continue
    return None

@router.post("/preview")
async def preview_excel(
    file: UploadFile = File(...),
    sheet_type: str = Form("activities"),
    current_user: models.User = Depends(get_current_user)
):
    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read Excel file: {str(e)}")

    if not all_rows:
        raise HTTPException(status_code=400, detail="Excel file is empty")

    # Find the header row (first non-empty row)
    header_idx = 0
    while header_idx < len(all_rows) and all(v is None for v in all_rows[header_idx]):
        header_idx += 1

    if header_idx >= len(all_rows):
        raise HTTPException(status_code=400, detail="No headers found in Excel file")

    raw_headers = all_rows[header_idx]
    columns = [str(h).strip() for h in raw_headers if h is not None]
    
    mapping = {}
    if sheet_type == "activities":
        expected_cols = {
            "faculty_name": ["faculty name", "faculty", "name", "teacher name"],
            "event_type": ["event details", "event type", "activity type", "type"],
            "organization": ["organization name", "organization", "institute", "host"],
            "title": ["title", "topic", "program title", "event title"],
            "month": ["month"],
            "start_date": ["start date", "from date", "start"],
            "end_date": ["end date", "to date", "end"],
            "certificate_link": ["certificates", "certificate", "certificate link", "proof"],
            "notes": ["any other notes", "notes", "remarks", "other notes"]
        }
    else:
        expected_cols = {
            "faculty_name": ["faculty name", "faculty", "name", "author"],
            "publication_type": ["publication", "publication type", "type"],
            "publication_details": ["publication details", "details", "journal/publisher"],
            "title": ["title", "paper title", "book title"],
            "month": ["month"],
            "date_of_publication": ["date of publication", "publication date", "date"],
            "milestone_achieved": ["milestone achieved", "milestone"],
            "certificate_link": ["certificates", "certificate", "link", "proof"],
            "notes": ["any other notes", "notes", "remarks"]
        }

    for target_key, aliases in expected_cols.items():
        for col in columns:
            if col.lower() in aliases:
                mapping[target_key] = col
                break

    data_rows = all_rows[header_idx + 1:]
    preview_rows = []
    for r in data_rows[:10]:
        if any(v is not None for v in r):
            row_dict = {}
            for col_idx, col_name in enumerate(columns):
                val = r[col_idx] if col_idx < len(r) else ""
                row_dict[col_name] = str(val) if val is not None else ""
            preview_rows.append(row_dict)

    return {
        "filename": file.filename,
        "total_rows": len([r for r in data_rows if any(v is not None for v in r)]),
        "columns": columns,
        "suggested_mapping": mapping,
        "preview_rows": preview_rows
    }

@router.post("/import")
async def import_excel(
    file: UploadFile = File(...),
    sheet_type: str = Form("activities"),
    column_mapping: str = Form("{}"),
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    import json
    try:
        col_map = json.loads(column_mapping)
    except Exception:
        col_map = {}

    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read Excel file: {str(e)}")

    if not all_rows:
        raise HTTPException(status_code=400, detail="Excel file is empty")

    header_idx = 0
    while header_idx < len(all_rows) and all(v is None for v in all_rows[header_idx]):
        header_idx += 1

    headers = [str(h).strip() if h is not None else f"Col{i}" for i, h in enumerate(all_rows[header_idx])]
    col_index_map = {h.lower(): i for i, h in enumerate(headers)}

    data_rows = all_rows[header_idx + 1:]
    current_fac = get_user_faculty(current_user, db)
    is_faculty = (current_user.role == models.RoleEnum.faculty.value)

    imported_count = 0
    errors = []

    for idx, row in enumerate(data_rows):
        if not any(v is not None for v in row):
            continue
        row_num = header_idx + idx + 2
        try:
            def get_val(mapped_name, default_col_name):
                target_col = col_map.get(mapped_name, default_col_name)
                c_idx = col_index_map.get(str(target_col).lower())
                if c_idx is not None and c_idx < len(row):
                    return row[c_idx]
                return None

            # Resolve Faculty
            faculty_id = None
            if is_faculty and current_fac:
                faculty_id = current_fac.id
            else:
                fac_name_val = str(get_val("faculty_name", "Faculty Name") or "").strip()
                if fac_name_val:
                    matched_fac = db.query(models.Faculty).filter(
                        (models.Faculty.name.ilike(f"%{fac_name_val}%")) |
                        (models.Faculty.employee_code == fac_name_val)
                    ).first()
                    if matched_fac:
                        faculty_id = matched_fac.id

            if not faculty_id:
                faculty_id = current_fac.id if current_fac else 1

            if sheet_type == "activities":
                title_val = str(get_val("title", "Title") or "").strip()
                if not title_val or title_val.lower() in ["title", "topic"]:
                    errors.append(f"Row {row_num}: Missing title")
                    continue

                event_type = str(get_val("event_type", "Event Details") or "Other").strip()
                organization = str(get_val("organization", "Organization Name") or "University").strip()
                month = str(get_val("month", "Month") or "").strip() or None
                start_date = parse_date_safely(get_val("start_date", "Start Date"))
                end_date = parse_date_safely(get_val("end_date", "End Date"))
                cert_link = str(get_val("certificate_link", "Certificates") or "").strip() or None
                notes = str(get_val("notes", "Any Other Notes") or "").strip() or None

                act = models.Activity(
                    faculty_id=faculty_id,
                    title=title_val,
                    event_type=event_type,
                    organization=organization,
                    month=month,
                    start_date=start_date,
                    end_date=end_date,
                    certificate_link=cert_link,
                    notes=notes,
                    status=models.StatusEnum.submitted.value if is_faculty else models.StatusEnum.approved.value,
                    created_at=datetime.datetime.utcnow()
                )
                db.add(act)
                imported_count += 1

            else: # publications
                title_val = str(get_val("title", "Title") or "").strip()
                if not title_val:
                    errors.append(f"Row {row_num}: Missing title")
                    continue

                pub_type = str(get_val("publication_type", "Publication") or "Book").strip()
                details = str(get_val("publication_details", "Publication Details") or "").strip() or None
                month = str(get_val("month", "Month") or "").strip() or None
                pub_date = parse_date_safely(get_val("date_of_publication", "Date of Publication"))
                milestone = str(get_val("milestone_achieved", "Milestone Achieved") or "Book Published").strip()
                cert_link = str(get_val("certificate_link", "Certificates") or "").strip() or None
                notes = str(get_val("notes", "Any Other Notes") or "").strip() or None

                pub = models.Publication(
                    faculty_id=faculty_id,
                    title=title_val,
                    publication_type=pub_type,
                    publication_details=details,
                    month=month,
                    date_of_publication=pub_date,
                    milestone_achieved=milestone,
                    certificate_link=cert_link,
                    notes=notes,
                    status=models.StatusEnum.submitted.value if is_faculty else models.StatusEnum.approved.value,
                    created_at=datetime.datetime.utcnow()
                )
                db.add(pub)
                imported_count += 1

        except Exception as e:
            errors.append(f"Row {row_num}: {str(e)}")

    db.commit()

    log_action(
        db=db,
        action="IMPORT",
        entity="Excel",
        user_id=current_user.id,
        user_email=current_user.email,
        details=f"Imported {imported_count} {sheet_type} records from Excel ({file.filename})"
    )

    return {
        "success": True,
        "imported_count": imported_count,
        "error_count": len(errors),
        "errors": errors[:20]
    }

@router.get("/export")
def export_excel(
    entity_type: str = "activities",
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    from app.services.reports import generate_excel_report

    fac = get_user_faculty(current_user, db)
    is_faculty = (current_user.role == models.RoleEnum.faculty.value)

    if entity_type == "activities":
        query = db.query(models.Activity).join(models.Faculty)
        if is_faculty and fac:
            query = query.filter(models.Activity.faculty_id == fac.id)
        records = query.order_by(models.Activity.created_at.desc()).all()

        headers = [
            "Faculty Name", "Event Details", "Organization Name", "Title",
            "Month", "Start Date", "End Date", "Certificates", "Any Other Notes", "Status"
        ]
        data = []
        for r in records:
            data.append([
                r.faculty.name if r.faculty else "",
                r.event_type or "",
                r.organization or "",
                r.title or "",
                r.month or "",
                str(r.start_date or ""),
                str(r.end_date or ""),
                r.certificate_link or "",
                r.notes or "",
                r.status or ""
            ])
        title = "Faculty Activities Export"
    else:
        query = db.query(models.Publication).join(models.Faculty)
        if is_faculty and fac:
            query = query.filter(models.Publication.faculty_id == fac.id)
        records = query.order_by(models.Publication.created_at.desc()).all()

        headers = [
            "Faculty Name", "Publication", "Publication Details", "Title",
            "Status", "Month", "Date of Publication", "Milestone Achieved", "Certificates", "Any Other Notes"
        ]
        data = []
        for r in records:
            data.append([
                r.faculty.name if r.faculty else "",
                r.publication_type or "Book",
                r.publication_details or "",
                r.title or "",
                r.status or "",
                r.month or "",
                str(r.date_of_publication or ""),
                r.milestone_achieved or "",
                r.certificate_link or "",
                r.notes or ""
            ])
        title = "Faculty Publications Export"

    log_action(
        db=db,
        action="EXPORT",
        entity="Excel",
        user_id=current_user.id,
        user_email=current_user.email,
        details=f"Exported {len(data)} {entity_type} records to Excel"
    )

    buf = generate_excel_report(title, headers, data, sheet_name=entity_type.capitalize())
    filename = f"{entity_type}_export_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
